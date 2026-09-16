#!/usr/bin/env python3
"""Build docs/index.html from data/raw_edits.json.

Derives the chart aggregates, builds the detail workbook, embeds it as base64, and
fills the template. Everything the dashboard needs ends up in the one HTML file, so
it works offline and has no second file to keep alongside it.
"""
import base64
import calendar
import datetime as dt
import difflib
import io
import json
import os
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, 'data', 'raw_edits.json')
TEMPLATE = os.path.join(ROOT, 'templates', 'dashboard.html')
OUT = os.path.join(ROOT, 'docs', 'index.html')
XLSX_NAME = 'mapped_sku_description_edits.xlsx'

# Workbook column order. Names that are not DERIVED are looked up in the extract by
# name, so reordering the SELECT in extract.sh cannot shift a column silently.
HEADERS = ['store_country', 'id_supplier_stock', 'fk_supplier', 'supplier_key', 'Supplier',
           'supplier_pid', 'fk_product_map', 'spinach_pid', 'previous_description',
           'new_description', 'Removed', 'Added', 'Change', 'previous_seen_at', 'changed_at',
           'Month']
DERIVED = {'Supplier', 'Removed', 'Added', 'Change', 'Month'}
WIDTHS = {'store_country': 13, 'id_supplier_stock': 17, 'fk_supplier': 12, 'supplier_key': 22,
          'Supplier': 18, 'supplier_pid': 18, 'fk_product_map': 15, 'spinach_pid': 15,
          'previous_description': 44, 'new_description': 44, 'Removed': 20, 'Added': 20,
          'Change': 30, 'previous_seen_at': 19, 'changed_at': 19, 'Month': 10}
INTS = {'id_supplier_stock', 'fk_supplier', 'fk_product_map'}
MONO_COLS = {'Removed', 'Added', 'Change'}


def parent(key):
    """Supplier name = supplier_key up to the first '-' or '_'."""
    for i, ch in enumerate(key):
        if ch in '-_':
            return key[:i]
    return key


def word_diff(prev, new):
    """Word-level diff: (removed, added, compact).

    Either side may be NULL: the extract only guarantees previous_description is
    set, so a description cleared to NULL arrives here as None.
    """
    a, b = (prev or '').split(), (new or '').split()
    rem, add = [], []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b).get_opcodes():
        if tag in ('replace', 'delete'):
            rem += a[i1:i2]
        if tag in ('replace', 'insert'):
            add += b[j1:j2]
    r, x = ' '.join(rem), ' '.join(add)
    if r and x:
        return r, x, f'{r} \u2192 {x}'
    if x:
        return r, x, f'+ {x}'
    if r:
        return r, x, f'\u2212 {r}'
    return r, x, 'reordered / spacing only'


MONTH_NAMES = ['January', 'February', 'March', 'April', 'May', 'June', 'July',
               'August', 'September', 'October', 'November', 'December']


def pretty(iso, with_year=True):
    """'2026-09-09' -> '9 September 2026'."""
    d = dt.date.fromisoformat(iso[:10])
    out = f'{d.day} {MONTH_NAMES[d.month - 1]}'
    return f'{out} {d.year}' if with_year else out


def ordinal(n):
    suffix = 'th' if 11 <= n % 100 <= 13 else {1: 'st', 2: 'nd', 3: 'rd'}.get(n % 10, 'th')
    return f'{n}{suffix}'


def month_span(first, last):
    """('2026-01', '2026-08') -> 'January to August'; same month -> 'January'."""
    def one(m):
        y, mo = m.split('-')
        return MONTH_NAMES[int(mo) - 1], y
    (fn, fy), (ln, ly) = one(first), one(last)
    if first == last:
        return fn if fy == ly else f'{fn} {fy}'
    if fy == ly:
        return f'{fn} to {ln}'
    return f'{fn} {fy} to {ln} {ly}'


def build_workbook(rows, I):
    """Detail workbook: one row per edit, with derived Supplier and diff columns."""
    missing = [h for h in HEADERS if h not in DERIVED and h not in I]
    if missing:
        sys.exit(f'Extract has no column named {", ".join(missing)} — '
                 'the workbook and scripts/extract.sh have drifted apart.')

    ARIAL, MONO = 'Arial', 'Consolas'
    hdr_font = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
    hdr_fill = PatternFill('solid', fgColor='1F3864')
    body, mono = Font(name=ARIAL, size=10), Font(name=MONO, size=10)
    thin = Border(bottom=Side(style='thin', color='D9D9D9'))

    skus = len({r[I['spinach_pid']] for r in rows})
    wb = Workbook()
    ws = wb.active
    ws.title = 'Edits'
    ws['A1'] = 'Mapped SKUs — every description edit'
    ws['A1'].font = Font(name=ARIAL, bold=True, size=12)
    ws['A2'] = (f'{len(rows):,} edits across {skus:,} distinct SKUs. Rows with no product mapping, '
                'and first-time population of the column, are excluded.')
    ws['A2'].font = Font(name=ARIAL, size=9, italic=True, color='595959')

    # Column number by name — nothing below counts positions by hand.
    COL = {h: i + 1 for i, h in enumerate(HEADERS)}

    hr = 4
    for h in HEADERS:
        cell = ws.cell(row=hr, column=COL[h])
        cell.value, cell.font, cell.fill = h, hdr_font, hdr_fill
        cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[hr].height = 20

    r_ = hr
    for r in rows:
        r_ += 1
        rem, add, compact = word_diff(r[I['previous_description']], r[I['new_description']])
        # Supplier and Month are written as values, not formulas: one implementation of
        # each rule, and the same one the charts aggregate on. Formulas written by
        # openpyxl carry no cached result, so pandas and most preview panes read them
        # as empty.
        derived = {
            'Supplier': parent(r[I['supplier_key']]),
            'Removed': rem,
            'Added': add,
            'Change': compact,
            'Month': r[I['changed_at']][:7],
        }
        for h in HEADERS:
            if h in derived:
                v = derived[h]
            else:
                v = r[I[h]]
                if v is not None and h in INTS:
                    v = int(v)
            cell = ws.cell(row=r_, column=COL[h])
            cell.value = v
            cell.font = mono if h in MONO_COLS else body
            cell.border = thin
            if h in INTS:
                cell.number_format = '0'
        ws.cell(row=r_, column=COL['Month']).alignment = Alignment(horizontal='center')

    for h in HEADERS:
        ws.column_dimensions[get_column_letter(COL[h])].width = WIDTHS[h]
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A{hr}:{get_column_letter(len(HEADERS))}{r_}'

    nt = wb.create_sheet('Notes')
    nt['A1'] = 'Source and scope'
    nt['A1'].font = Font(name=ARIAL, bold=True, size=12)
    notes = [
        ('Scope', 'fk_product_map IS NOT NULL AND previous_description IS NOT NULL'),
        ('Supplier', 'Derived value: supplier_key up to the first "-" or "_". Groups a supplier\'s state '
                     'and branch feeds together. Also folds in non-location variants (promo, '
                     'virtual, special) — keep supplier_key when those must stay separate.'),
        ('Removed / Added / Change', 'Derived by comparing previous_description with '
                                     'new_description word by word. "Change" is the compact form. '
                                     '"reordered / spacing only" means the same words in a '
                                     'different order or spacing.'),
        ('One edit, several rows', 'Each state-level account holds its own id_supplier_stock for '
                                   'the same product, so one upstream edit appears once per '
                                   'account. Count distinct spinach_pid for SKUs, row count for '
                                   'stock records.'),
        ('Mapping is point-in-time', 'fk_product_map is the mapping as it stood when the change '
                                     'was captured, not today.'),
        ('Supplier identity', 'Neither supplier_key nor fk_supplier is unique on its own — group '
                              'on the pair.'),
        ('Timestamps', 'UTC, written as text so they survive the export unchanged.'),
    ]
    r_ = 2
    for label, val in notes:
        r_ += 1
        nt.cell(row=r_, column=1).value = label
        nt.cell(row=r_, column=1).font = Font(name=ARIAL, size=10, bold=True)
        nt.cell(row=r_, column=1).alignment = Alignment(vertical='top')
        nt.cell(row=r_, column=2).value = val
        nt.cell(row=r_, column=2).font = body
        nt.cell(row=r_, column=2).alignment = Alignment(wrap_text=True, vertical='top')
        nt.row_dimensions[r_].height = 40
    nt.column_dimensions['A'].width = 26
    nt.column_dimensions['B'].width = 95

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main():
    if not os.path.exists(RAW):
        sys.exit(f'Missing {RAW} — run scripts/extract.sh first.')
    d = json.load(open(RAW))
    cols, rows = d['cols'], d['rows']
    I = {c: i for i, c in enumerate(cols)}

    # Columns are addressed by name throughout, so a reordered SELECT is harmless and a
    # renamed one stops the build here rather than shifting data into the wrong column.
    missing = [h for h in HEADERS if h not in DERIVED and h not in I]
    if missing:
        sys.exit(f'Extract has no column named {", ".join(missing)} — scripts/extract.sh '
                 f'and build.py have drifted apart. Got: {", ".join(cols)}')

    SP, SK, CH = I['spinach_pid'], I['supplier_key'], I['changed_at']

    if not rows:
        sys.exit('Extract is empty — nothing to build.')

    # --- aggregates: distinct SKUs, not rows ---
    psk, ped = defaultdict(set), Counter()
    for r in rows:
        p = parent(r[SK])
        psk[p].add(r[SP])
        ped[p] += 1
    suppliers = sorted(({'p': p, 'skus': len(psk[p]), 'edits': ped[p]} for p in psk),
                       key=lambda x: -x['skus'])

    msk, med = defaultdict(set), Counter()
    for r in rows:
        m = r[CH][:7]
        msk[m].add(r[SP])
        med[m] += 1
    months = [{'m': m, 'skus': len(msk[m]), 'edits': med[m]} for m in sorted(msk)]

    stats = {'skus': len({r[SP] for r in rows}), 'edits': len(rows), 'suppliers': len(suppliers)}

    # --- everything the page used to hardcode, derived from the data ---
    first_iso, last_iso = min(r[CH] for r in rows), max(r[CH] for r in rows)
    last_day = dt.date.fromisoformat(last_iso[:10])
    eom = calendar.monthrange(last_day.year, last_day.month)[1]

    # The final month is partial unless the data reaches its last calendar day.
    partial = len(months) - 1 if last_day.day < eom else -1
    if partial >= 0:
        pm = months[partial]['m']
        first_day = min(dt.date.fromisoformat(r[CH][:10]).day for r in rows if r[CH][:7] == pm)
        name = MONTH_NAMES[int(pm.split('-')[1]) - 1]
        partial_note = f'{first_day} to {last_day.day} {name} only'
        partial_tick = f'to the {ordinal(last_day.day)}'
    else:
        partial_note = partial_tick = ''

    # The average line covers whole months only, and needs at least two to mean anything.
    done = [m for i, m in enumerate(months) if i != partial]
    avg = round(sum(m['skus'] for m in done) / len(done)) if len(done) >= 2 else None
    avg_range = month_span(done[0]['m'], done[-1]['m']) if avg is not None else ''

    extracted = d.get('extracted') or dt.date.fromtimestamp(os.path.getmtime(RAW)).isoformat()
    same_year = first_iso[:4] == last_iso[:4]
    date_range = (f'between {pretty(first_iso, with_year=not same_year)} '
                  f'and {pretty(last_iso)}')

    xlsx = build_workbook(rows, I)

    payload = {
        'suppliers': suppliers,
        'months': months,
        'stats': stats,
        'meta': {
            'dateRange': date_range,
            'extracted': pretty(extracted),
            'partial': partial,
            'partialNote': partial_note,
            'partialTick': partial_tick,
            'avg': avg,
            'avgRange': avg_range,
            'columns': len(HEADERS),
            'xlsxKb': round(len(xlsx) / 1e3),
        },
    }

    tpl = open(TEMPLATE).read()
    for token in ('__DATA__', '__XLSX__', '__FILE__'):
        if token not in tpl:
            sys.exit(f'Template is missing the {token} placeholder.')
    data = json.dumps(payload, separators=(',', ':')).replace('<', '\\u003c').replace('&', '\\u0026')
    html = (tpl.replace('__DATA__', data)
               .replace('__XLSX__', base64.b64encode(xlsx).decode())
               .replace('__FILE__', XLSX_NAME))
    for token in ('__DATA__', '__XLSX__', '__FILE__'):
        if token in html:
            sys.exit(f'{token} survived substitution — the build would ship a broken page.')

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    open(OUT, 'w').write(html)

    print(f'{stats["edits"]:,} edits · {stats["skus"]:,} SKUs · {stats["suppliers"]} suppliers')
    print(f'{date_range}, extracted {pretty(extracted)}')
    if partial >= 0:
        print(f'last month {months[partial]["m"]} drawn as partial ({partial_note})')
    else:
        print(f'last month {months[-1]["m"]} is complete — no partial bar')
    print(f'average line: {avg if avg is not None else "suppressed"}'
          f'{" a month, " + avg_range if avg is not None else " (fewer than two whole months)"}')
    print(f'{OUT} — {os.path.getsize(OUT) / 1e6:.2f} MB, workbook {len(xlsx) / 1e3:.0f} KB embedded')


if __name__ == '__main__':
    main()
